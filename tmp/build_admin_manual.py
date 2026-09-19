from __future__ import annotations

import os
from datetime import datetime
from pathlib import Path

from PIL import Image as PILImage
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    ListFlowable,
    ListItem,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables"
ASSETS = OUT / "assets"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)

DOCX_PATH = OUT / "Matrix-Intertrade-Admin-Content-Management-Manual.docx"
PDF_PATH = OUT / "Matrix-Intertrade-Admin-Content-Management-Manual.pdf"

TEMP = Path(r"C:\Users\1137~1\AppData\Local\Temp")
SCREENSHOTS = {
    "signin": ASSETS / "admin-signin.png",
    "tabs": TEMP / "codex-clipboard-ea779b19-4b74-44bb-a6de-ca2a7a956aec.png",
    "actions_product": TEMP / "codex-clipboard-0f26f122-9858-4564-a17d-6953ea7d5779.png",
    "actions_settings": TEMP / "codex-clipboard-9faf9dcf-6766-460f-8177-a88b0986cffd.png",
    "product_public": TEMP / "codex-clipboard-92d14bda-c01f-4448-bf0a-2c4329bf819e.png",
    "product_editor": TEMP / "codex-clipboard-dbf63683-7b71-41f9-9492-4f9f19862cf0.png",
    "rich_text": TEMP / "codex-clipboard-4f59e886-3a03-4eb3-8cb6-f8c87e3f8e36.png",
    "industry_image_hidden": ASSETS / "industry-image-hidden-redacted.png",
}

NAVY = "001B44"
BLUE = "008AF5"
CYAN = "02BDEB"
PALE = "EEF7FF"
MUTED = "5B677A"
GOLD = "F4B000"
RED = "EF4444"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False, color: str = NAVY) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Tahoma"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_run(run, size=10.5, color=NAVY, bold=False, italic=False) -> None:
    run.font.name = "Tahoma"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_doc_para(doc, text: str, style: str | None = None, *, bold=False, italic=False, size=10.5, color=NAVY, after=6, align=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_doc_bullets(doc, items: list[str], level: int = 0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run(r, size=9.8, color=NAVY)


def add_doc_heading(doc, text: str, level: int = 1) -> None:
    style = f"Heading {level}"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run(r, size=16 if level == 1 else 13 if level == 2 else 11.5, color=NAVY if level == 1 else BLUE, bold=True)


def add_doc_table(doc, rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(rows[0]):
        set_cell_text(hdr[i], h, bold=True, color=NAVY)
        set_cell_shading(hdr[i], "DFF3FF")
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_text(cells[i], text)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()


def fit_image(path: Path, max_width_in=6.15, max_height_in=3.3) -> tuple[float, float] | None:
    if not path.exists():
        return None
    with PILImage.open(path) as img:
        w, h = img.size
    scale = min(max_width_in / (w / 96), max_height_in / (h / 96), 1.0)
    return (w / 96 * scale, h / 96 * scale)


def add_doc_image(doc, key: str, caption: str, max_width=6.15, max_height=3.3) -> None:
    path = SCREENSHOTS.get(key)
    dims = fit_image(path, max_width, max_height) if path else None
    if not dims:
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(dims[0]), height=Inches(dims[1]))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    set_run(r, size=8.5, color=MUTED, italic=True)
    cap.paragraph_format.space_after = Pt(8)


def setup_doc_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    for style_name in ["Normal", "List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Tahoma"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
        style.font.size = Pt(10.5)
    for i, (size, color) in enumerate([(16, NAVY), (13, BLUE), (11.5, NAVY)], start=1):
        style = doc.styles[f"Heading {i}"]
        style.font.name = "Tahoma"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Tahoma")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
    header = section.header.paragraphs[0]
    header.text = "Matrix Intertrade | Admin Content Management Manual"
    header.runs[0].font.name = "Tahoma"
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    footer = section.footer.paragraphs[0]
    footer.text = "คู่มือสำหรับลูกค้า - ใช้ประกอบการดูแลเว็บไซต์"
    footer.runs[0].font.name = "Tahoma"
    footer.runs[0].font.size = Pt(8.5)
    footer.runs[0].font.color.rgb = RGBColor.from_string(MUTED)


MENU_ROWS = [
    ["เมนู", "ใช้จัดการ", "แสดงผลบนเว็บไซต์"],
    ["All Products", "สินค้า รูปภาพ ราคา แบรนด์ รายละเอียด Rich Text และ SEO", "หน้ารวมสินค้า, filter แบรนด์, หน้ารายละเอียดสินค้า"],
    ["Articles", "บทความ รูปปก หมวดหมู่ เนื้อหา สถานะเผยแพร่ และ SEO", "หน้า Blog, หน้ารายละเอียดบทความ, Footer บทความ"],
    ["Article Categories", "หมวดบทความและรูปหมวด", "ตัวกรอง/กลุ่มบทความในหน้า Blog"],
    ["Brands", "ข้อมูลแบรนด์ โลโก้ รูปภาพ หน้าแบรนด์ และ SEO", "หน้า Brands, หน้า /brands/{slug}, หน้า category/filter"],
    ["Solutions", "บริการหลัก เช่น จอ LED, Interactive Display, AV Solutions", "หน้าแรก, เมนูโซลูชั่น, หน้า solution detail"],
    ["Industries", "กลุ่มลูกค้า showcase card, metric, หน้า detail และ SEO", "หน้าแรก, /brands, /aboutus, /industry/{slug}"],
    ["เกี่ยวกับเรา (About)", "ประวัติบริษัท, story, mission, vision, values, stats", "หน้า /aboutus"],
    ["Contact Submissions", "รายการ lead จากแบบฟอร์มติดต่อและสถานะอ่านแล้ว", "ข้อมูลหลังบ้านเท่านั้น"],
    ["Contact Us", "ข้อความหน้า contact, เบอร์, อีเมล, LINE, map และ SEO", "หน้า /contactus"],
    ["Footer", "CTA, company description, contact info, social links, newsletter", "Footer ทุกหน้า"],
    ["LINE Settings", "Token/ปลายทางแจ้งเตือน LINE สำหรับ contact form", "ระบบแจ้งเตือนหลังส่งฟอร์ม"],
    ["Tracking Settings", "Google Analytics / Meta Pixel", "Tracking script ฝั่ง public website"],
]

SECTIONS = [
    (
        "All Products",
        "ใช้จัดการข้อมูลสินค้าทั้งหมดที่แสดงบนเว็บไซต์ ตั้งแต่ชื่อสินค้า รูปภาพ ราคา แบรนด์ ไปจนถึงรายละเอียดสินค้าแบบ Rich Text และข้อมูล SEO.",
        [
            "Product ID เป็นรหัสอ้างอิงสินค้า ไม่ควรแก้เองหากระบบสร้างให้แล้ว",
            "Name และ Slug ส่งผลต่อชื่อหน้าและ URL สินค้า",
            "Description Text ใช้เป็นคำอธิบายสั้นในหน้ารวมสินค้าและ fallback สำหรับ SEO",
            "Description HTML หรือ Rich Text ใช้เป็นรายละเอียดหลักบนหน้าสินค้า",
            "Brand / Brand Slug / Category ID ต้องสอดคล้องกันเพื่อให้ filter และ breadcrumb ถูกต้อง",
            "Price Text และ Source URL แสดงในหน้ารายละเอียดสินค้าและช่วยอ้างอิงข้อมูล",
            "SEO/Social ใช้ควบคุม title, description, image และ no-index",
        ],
        "Products/Articles ใช้ workflow แบบ Save draft และ Publish. Draft จะเก็บข้อมูลไว้ก่อน ส่วน Publish คือข้อมูลที่พร้อมแสดงหน้าบ้าน.",
    ),
    (
        "Articles",
        "ใช้สร้างและแก้ไขบทความ ข่าวสาร คู่มือ และ content SEO ที่แสดงในหน้า Blog.",
        [
            "Title และ Slug กำหนดชื่อบทความและ URL",
            "Category ควรเลือกจาก Article Categories เพื่อให้จัดหมวดถูกต้อง",
            "Cover Image แสดงในหน้ารวม Blog และด้านบนบทความ",
            "Excerpt เป็นข้อความสรุปสั้นสำหรับ preview และ SEO fallback",
            "Content HTML/Rich Text เป็นเนื้อหาบทความจริง",
            "Featured และ Publish Date ใช้กำหนดความเด่นและเวลาที่แสดง",
            "SEO fields ช่วยควบคุมผลลัพธ์ Google และ Social share",
        ],
        "ก่อน Publish ควรตรวจรูปปก, paragraph แรก, heading, link และ preview ใน Google Search Preview.",
    ),
    (
        "Article Categories",
        "ใช้จัดการหมวดหมู่บทความ เช่น LED Display, Interactive Display, Smart Classroom หรือ Case Study.",
        [
            "Slug ควรเป็นภาษาอังกฤษ ตัวเล็ก และไม่ซ้ำ",
            "Label คือชื่อหมวดที่ลูกค้าเห็น",
            "Image URL ใช้เป็นภาพประกอบหมวด ถ้ามี",
            "Payload JSON ใช้เฉพาะกรณีที่ทีมเทคนิคกำหนดข้อมูลเสริม",
        ],
        "ไม่ควรลบหมวดที่มีบทความใช้งานอยู่ เว้นแต่ย้ายบทความไปหมวดใหม่ก่อน.",
    ),
    (
        "Brands",
        "ใช้จัดการหน้าแบรนด์ รูปโลโก้ ข้อมูลหมวดแบรนด์ และข้อมูล SEO ของแบรนด์.",
        [
            "Slug เป็น URL ของหน้าแบรนด์ เช่น /brands/grandview",
            "Name และ Category แสดงบน card และหัวข้อหน้าแบรนด์",
            "Description ใช้อธิบายภาพรวมแบรนด์",
            "Image URL และ Logo URL ควรใช้ภาพคมชัด พื้นหลังเหมาะกับเว็บ",
            "Category Info เช่น tagline, intro, highlights ใช้กับหน้า category/brand detail",
            "SEO fields ช่วยให้หน้าแบรนด์ค้นหาเจอได้ดีขึ้น",
        ],
        "ระวังการแก้ slug เพราะจะทำให้ link เดิมเปลี่ยนและอาจกระทบ SEO.",
    ),
    (
        "Solutions",
        "ใช้จัดการบริการหลักของบริษัท เช่น จอ LED, Projection Screen, Wireless Presentation และ AV Solutions.",
        [
            "Title/Description ใช้ใน card และหน้า solution",
            "Image URL เป็นภาพหลักของ solution",
            "Payload JSON เก็บข้อมูลหน้า detail เช่น banner, FAQ, product sections หรือ CTA",
            "SEO fields ควรเขียนให้ตรง intent ของคำค้น เช่น จอ LED สำหรับองค์กร",
        ],
        "ควรแก้ Payload JSON เฉพาะเมื่อเข้าใจโครงสร้าง หรือให้ทีมเทคนิคช่วยตรวจ.",
    ),
    (
        "Industries",
        "ใช้จัดการกลุ่มลูกค้าที่บริษัทให้บริการ เช่น โรงเรียน โรงแรม องค์กร และระบบประชุมทางไกล.",
        [
            "Title/Description ใช้ทั้งหน้า industry detail และ card บนหน้าอื่น",
            "Showcase Card Image แสดงใน card บนหน้าแรก, /brands และ /aboutus",
            "Show on homepage and /brands ใช้ซ่อน/แสดง card",
            "Card Order กำหนดลำดับการแสดง",
            "Metric Value/Label ใช้กับตัวเลข highlight บน card",
            "Payload JSON ใช้กับรายละเอียดหน้า /industry/{slug}",
            "Image URL เดิมถูกซ่อนเพื่อลดความสับสนในการใช้งาน",
        ],
        "หลังแก้ Industries ควรตรวจ 4 จุด: หน้าแรก, /brands, /aboutus และ /industry/{slug}.",
    ),
    (
        "เกี่ยวกับเรา (About)",
        "ใช้จัดการข้อมูลบริษัทและเนื้อหาหน้า /aboutus ที่ลูกค้าเห็นจริง.",
        [
            "Story P1/P2/P3 คือย่อหน้าประวัติและแนวคิดบริษัท",
            "Stats Payload คือชุดตัวเลข เช่น 20+, 500+, 8+, 100%",
            "Mission, Vision, Values แสดงในการ์ดบนหน้าเกี่ยวกับเรา",
            "ส่วน (EN) ถูกตัดออกเพื่อลดความซับซ้อนตาม requirement ล่าสุด",
        ],
        "เหมาะให้ admin ทั่วไปแก้ข้อความหลักได้ แต่ควรตรวจความยาวเพื่อไม่ให้ layout ยาวเกินไป.",
    ),
    (
        "Contact Submissions",
        "ใช้ดูรายการติดต่อจากลูกค้าที่ส่งผ่านฟอร์มหน้า /contactus.",
        [
            "ดูชื่อ เบอร์ อีเมล หัวข้อ และข้อความที่ลูกค้าส่งมา",
            "ใช้สถานะ read/unread เพื่อติดตามว่าทีมตอบกลับแล้วหรือยัง",
            "ไม่ควรแก้ข้อความ lead เดิม เพราะเป็นข้อมูลอ้างอิงจากลูกค้า",
        ],
        "ข้อมูลในเมนูนี้เป็นข้อมูลภายใน ไม่แสดงหน้าบ้าน.",
    ),
    (
        "Contact Us",
        "ใช้จัดการข้อมูลหลักในหน้า /contactus เช่น hero, ข้อมูลติดต่อ, map และ SEO.",
        [
            "Hero Title/Description คือข้อความนำบนหน้า contact",
            "Phone, Email, LINE และ Address ใช้แสดงช่องทางติดต่อ",
            "Map Embed URL ใช้แสดงแผนที่",
            "Directions URL ใช้กับปุ่มนำทาง Google Maps",
            "SEO/JSON-LD ช่วยให้ Google เข้าใจข้อมูลบริษัทและที่อยู่",
        ],
        "ก่อน save ควรทดสอบเบอร์โทร, อีเมล และปุ่มนำทางว่าเปิดได้ถูกต้อง.",
    ),
    (
        "Footer",
        "ใช้จัดการข้อมูลส่วนท้ายเว็บที่แสดงทุกหน้า เช่น CTA, company description, contact info, social links และ newsletter.",
        [
            "CTA Strip คือข้อความเชิญชวนและปุ่มด้านบน footer",
            "Company Description คือคำอธิบายบริษัทด้านซ้าย",
            "Address/Phone/Email/LINE ควรตรงกับหน้า Contact Us",
            "Social Links ใช้กับไอคอน social",
            "Newsletter Text คือข้อความในกล่องสมัครรับข่าวสาร",
        ],
        "รายการ Solutions, Brands และ Articles ใน footer ยังใช้ logic เดิม ไม่ได้แก้จากเมนูนี้ในรอบนี้.",
    ),
    (
        "LINE Settings",
        "ใช้ตั้งค่าการแจ้งเตือน LINE เมื่อมีลูกค้าส่งฟอร์มติดต่อ.",
        [
            "Token หรือ Credential เป็นข้อมูลลับ ห้ามส่งต่อในเอกสารหรือแชตทั่วไป",
            "ใช้ปุ่ม test เฉพาะตอนตั้งค่าหรือแก้ไขปลายทางแจ้งเตือน",
            "ถ้าแจ้งเตือนไม่เข้า ให้ตรวจ token, ปลายทาง, quota และ log",
        ],
        "ควรให้ทีมเทคนิคหรือผู้ดูแลระบบหลักเป็นคนแก้เมนูนี้.",
    ),
    (
        "Tracking Settings",
        "ใช้ตั้งค่า Google Analytics และ Meta Pixel สำหรับวัดผลบนเว็บไซต์.",
        [
            "Google Analytics ID ใช้สำหรับวัด traffic และ event",
            "Meta Pixel ID ใช้สำหรับ remarketing และ conversion",
            "ควรใส่ ID ที่ผ่านการตรวจจากบัญชีจริงเท่านั้น",
            "ไม่ควรใส่ code script เต็มถ้าระบบต้องการเฉพาะ ID",
        ],
        "ควรทดสอบผ่าน browser devtools หรือเครื่องมือของ Google/Meta หลัง deploy.",
    ),
]


def build_docx() -> None:
    doc = Document()
    setup_doc_styles(doc)
    add_doc_para(doc, "MATRIX INTERTRADE", size=10, color=BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=16)
    add_doc_para(doc, "คู่มือการใช้งานระบบหลังบ้าน", size=24, color=NAVY, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    add_doc_para(doc, "Admin Content Management Manual", size=15, color=BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    add_doc_para(doc, "สำหรับดูแลสินค้า บทความ แบรนด์ โซลูชั่น ข้อมูลบริษัท Contact Us, Footer, LINE และ Tracking Settings", size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=26)
    add_doc_para(doc, f"Version 1.0 | จัดทำเมื่อ {datetime.now().strftime('%d/%m/%Y')}", size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=20)
    add_doc_image(doc, "tabs", "ภาพรวมเมนู Content Management ในระบบหลังบ้าน", max_height=0.75)
    doc.add_page_break()

    add_doc_heading(doc, "สารบัญ", 1)
    toc_items = [
        "1. ภาพรวมระบบหลังบ้าน",
        "2. ความหมายของปุ่มสำคัญ",
        "3. คู่มือรายเมนู",
        "4. Workflow แนะนำก่อนเผยแพร่",
        "5. Checklist และ Troubleshooting",
        "6. ข้อควรระวังด้าน SEO, รูปภาพ, JSON และข้อมูลลับ",
    ]
    add_doc_bullets(doc, toc_items)
    doc.add_page_break()

    add_doc_heading(doc, "1. ภาพรวมระบบหลังบ้าน", 1)
    add_doc_para(doc, "เข้าระบบผ่าน /admin แล้วลงชื่อเข้าใช้ด้วยบัญชี Supabase admin ที่ได้รับสิทธิ์ ระบบนี้ใช้สำหรับแก้ไขข้อมูล runtime ที่แสดงบนเว็บไซต์ โดยหลังบันทึกข้อมูลแล้วควรเปิดหน้าจริงเพื่อตรวจผลเสมอ", color=NAVY)
    add_doc_image(doc, "signin", "หน้าล็อกอิน Admin จากระบบจริงล่าสุด", max_height=2.7)
    add_doc_table(doc, MENU_ROWS, widths=[1.35, 2.6, 2.4])

    add_doc_heading(doc, "2. ความหมายของปุ่มสำคัญ", 1)
    add_doc_bullets(doc, [
        "Refresh: โหลดข้อมูลล่าสุดจาก Supabase อีกครั้ง ใช้หลังมีการแก้ไขข้อมูลหรือสงสัยว่าหน้าจอยังค้าง state เก่า",
        "Last synced: เวลาล่าสุดที่ระบบโหลดข้อมูลสำเร็จ ช่วยยืนยันว่าข้อมูลบนหน้าจอเป็นชุดล่าสุด",
        "Save draft: บันทึกเป็นฉบับร่างสำหรับ Products/Articles โดยยังไม่จำเป็นต้องเผยแพร่",
        "Publish: เผยแพร่ Products/Articles ให้พร้อมแสดงหน้าบ้าน",
        "Save changes: บันทึกเมนูที่ไม่มีระบบ draft/publish เช่น Brands, Solutions, Industries, About และ Settings",
        "Discard changes: ยกเลิกการแก้ไขที่ยังไม่ได้บันทึกและย้อนกลับไปค่าล่าสุด",
        "Delete: ลบรายการออกจากระบบ ควรใช้ด้วยความระมัดระวัง",
    ])
    add_doc_image(doc, "actions_product", "ชุดปุ่ม Save draft / Publish / Delete สำหรับสินค้าและบทความ", max_height=1.25)
    add_doc_image(doc, "actions_settings", "ชุดปุ่ม Save changes / Discard changes / Delete สำหรับเมนูทั่วไป", max_height=1.8)

    add_doc_heading(doc, "3. คู่มือรายเมนู", 1)
    for title, intro, bullets, note in SECTIONS:
        add_doc_heading(doc, title, 2)
        add_doc_para(doc, intro, color=NAVY, after=4)
        add_doc_bullets(doc, bullets)
        add_doc_para(doc, f"ข้อแนะนำ: {note}", size=9.8, color=MUTED, italic=True)
        if title == "All Products":
            add_doc_image(doc, "product_editor", "ตัวอย่างช่องคำอธิบายสินค้าในระบบหลังบ้าน", max_height=1.45)
            add_doc_image(doc, "product_public", "ตัวอย่างข้อมูลสินค้าที่แสดงในหน้าบ้านหลังอัปเดต", max_height=3.6)
        if title == "Articles":
            add_doc_image(doc, "rich_text", "Rich Text editor มี scroll bar เพื่อรองรับเนื้อหายาว", max_height=3.1)
        if title == "Industries":
            add_doc_image(doc, "industry_image_hidden", "ตัวอย่างส่วน Image URL เดิมที่ถูกซ่อนเพื่อไม่ให้ admin สับสน", max_height=2.1)

    doc.add_page_break()
    add_doc_heading(doc, "4. Workflow แนะนำสำหรับลูกค้า", 1)
    add_doc_table(doc, [
        ["ขั้นตอน", "สิ่งที่ต้องทำ", "ผลลัพธ์ที่ควรตรวจ"],
        ["1", "เลือกเมนูและค้นหารายการ", "รายการที่เลือกตรงกับหน้าที่ต้องการแก้"],
        ["2", "แก้ข้อมูลในฟอร์ม", "Preview ด้านขวาอัปเดตตามข้อมูลใหม่"],
        ["3", "ตรวจ SEO/รูป/Slug", "ไม่มีข้อความซ้ำ รูปโหลดได้ URL ถูกต้อง"],
        ["4", "Save draft หรือ Save changes", "ไม่มี error และ Last synced อัปเดตหลัง Refresh"],
        ["5", "Publish ถ้าเป็น Products/Articles", "หน้าบ้านแสดงข้อมูลตามที่ตั้งใจ"],
        ["6", "เปิดหน้าจริงตรวจซ้ำ", "ไม่มี layout เพี้ยน รูปแตก หรือ link ผิด"],
    ], widths=[0.7, 3.0, 2.75])

    add_doc_heading(doc, "5. Checklist ก่อนเผยแพร่สินค้า/บทความ", 1)
    add_doc_bullets(doc, [
        "ชื่อเรื่องหรือชื่อสินค้าชัดเจนและสะกดถูก",
        "Slug เป็นภาษาอังกฤษ ตัวเล็ก ใช้ hyphen และไม่ซ้ำ",
        "รูปหลักคมชัด ไม่บีบผิดสัดส่วน และไม่มีข้อมูลลับ",
        "Description Text สรุปใจความ 1-3 ประโยค",
        "Rich Text มี heading, bullet, paragraph และ link ที่อ่านง่าย",
        "SEO Title และ Meta Description ไม่ยาวเกินไปและตรง keyword หลัก",
        "Preview / Google Search Preview ดูเรียบร้อย",
        "เปิด URL หน้าจริงหลังบันทึกเพื่อตรวจผล",
    ])

    add_doc_heading(doc, "6. Troubleshooting เบื้องต้น", 1)
    add_doc_table(doc, [
        ["ปัญหา", "สาเหตุที่พบบ่อย", "วิธีแก้"],
        ["Save ไม่ผ่าน", "Session หมดอายุ, JSON ผิดรูปแบบ, field จำเป็นว่าง", "Refresh, login ใหม่, ตรวจ error และ JSON"],
        ["รูปไม่แสดง", "URL ผิด, ไฟล์ไม่ public, ขนาดใหญ่เกินไป", "Upload ใหม่หรือใช้ URL ที่ browser เปิดได้"],
        ["หน้าบ้านไม่อัปเดต", "ยังไม่ได้ Publish, cache หรือยังไม่ Refresh", "กด Publish/Save, Refresh admin และ reload หน้าบ้าน"],
        ["Slug ซ้ำ", "ใช้ URL ซ้ำกับรายการเดิม", "เปลี่ยน slug ให้เฉพาะเจาะจง"],
        ["JSON error", "ลืม comma/quote หรือโครงสร้างผิด", "ตรวจด้วย JSON validator ก่อนบันทึก"],
        ["LINE ไม่แจ้งเตือน", "Token ผิด, quota, endpoint ปิด", "ตรวจ LINE Settings และทดสอบอีกครั้ง"],
    ], widths=[1.35, 2.45, 2.65])

    add_doc_heading(doc, "ข้อควรระวังด้านข้อมูลลับ", 1)
    add_doc_bullets(doc, [
        "ห้ามส่งต่อ LINE token, Supabase key, tracking ID แบบเต็ม หรือรหัสผ่าน admin ในเอกสารสาธารณะ",
        "ภาพหน้าจอสำหรับคู่มือควร crop เฉพาะ UI ที่ต้องอธิบาย",
        "เมนู LINE Settings และ Tracking Settings ควรจำกัดให้ผู้ดูแลระบบหลักหรือทีมเทคนิคแก้ไข",
        "ก่อนส่งคู่มือให้ลูกค้า ควรตรวจอีกครั้งว่าไม่มีค่า secret หรือข้อมูลส่วนตัวติดอยู่ในภาพ",
    ])

    doc.save(DOCX_PATH)


def register_fonts() -> str:
    regular = r"C:\Windows\Fonts\tahoma.ttf"
    bold = r"C:\Windows\Fonts\tahomabd.ttf"
    pdfmetrics.registerFont(TTFont("Tahoma", regular))
    pdfmetrics.registerFont(TTFont("Tahoma-Bold", bold))
    return "Tahoma"


def p(text: str, style):
    return Paragraph(text.replace("\n", "<br/>"), style)


def bullet_list(items: list[str], style):
    return ListFlowable(
        [ListItem(p(item, style), leftIndent=12) for item in items],
        bulletType="bullet",
        start="circle",
        leftIndent=18,
        bulletFontName="Tahoma",
        bulletFontSize=8,
    )


def img_flow(key: str, caption: str, max_w=6.55 * inch, max_h=3.35 * inch):
    path = SCREENSHOTS.get(key)
    if not path or not path.exists():
        return []
    with PILImage.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h, 1)
    return [
        Image(str(path), width=w * scale, height=h * scale),
        Spacer(1, 4),
        p(caption, STYLES["caption"]),
        Spacer(1, 10),
    ]


def table_flow(rows: list[list[str]], col_widths: list[float]):
    data = [[p(cell, STYLES["table_header" if r == 0 else "table_cell"]) for cell in row] for r, row in enumerate(rows)]
    table = Table(data, colWidths=[w * inch for w in col_widths], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DFF3FF")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#001B44")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D7E4EF")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 6),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return [table, Spacer(1, 12)]


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont("Tahoma", 7.5)
    canvas.setFillColor(colors.HexColor("#5B677A"))
    canvas.drawString(0.78 * inch, 0.42 * inch, "Matrix Intertrade - Admin Content Management Manual")
    canvas.drawRightString(7.72 * inch, 0.42 * inch, f"Page {doc.page}")
    canvas.restoreState()


def build_pdf() -> None:
    register_fonts()
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        rightMargin=0.72 * inch,
        leftMargin=0.72 * inch,
        topMargin=0.68 * inch,
        bottomMargin=0.68 * inch,
        title="Matrix Intertrade Admin Content Management Manual",
    )
    story = []
    story.append(Spacer(1, 1.05 * inch))
    story.append(p("MATRIX INTERTRADE", STYLES["kicker_center"]))
    story.append(p("คู่มือการใช้งานระบบหลังบ้าน", STYLES["cover_title"]))
    story.append(p("Admin Content Management Manual", STYLES["cover_subtitle"]))
    story.append(Spacer(1, 14))
    story.append(p("สำหรับดูแลสินค้า บทความ แบรนด์ โซลูชั่น ข้อมูลบริษัท Contact Us, Footer, LINE และ Tracking Settings", STYLES["lead_center"]))
    story.append(Spacer(1, 16))
    story.extend(img_flow("tabs", "ภาพรวมเมนู Content Management ในระบบหลังบ้าน", max_h=0.65 * inch))
    story.append(Spacer(1, 20))
    story.append(p(f"Version 1.0 | จัดทำเมื่อ {datetime.now().strftime('%d/%m/%Y')}", STYLES["muted_center"]))
    story.append(PageBreak())

    story.append(p("สารบัญ", STYLES["h1"]))
    story.append(bullet_list([
        "ภาพรวมระบบหลังบ้าน",
        "ความหมายของปุ่มสำคัญ",
        "คู่มือรายเมนูครบทุก tab",
        "Workflow แนะนำก่อนเผยแพร่",
        "Checklist และ Troubleshooting",
        "ข้อควรระวังด้าน SEO, รูปภาพ, JSON และข้อมูลลับ",
    ], STYLES["body"]))
    story.append(PageBreak())

    story.append(p("1. ภาพรวมระบบหลังบ้าน", STYLES["h1"]))
    story.append(p("เข้าระบบผ่าน /admin แล้วลงชื่อเข้าใช้ด้วยบัญชี Supabase admin ที่ได้รับสิทธิ์ ระบบนี้ใช้สำหรับแก้ไขข้อมูล runtime ที่แสดงบนเว็บไซต์ โดยหลังบันทึกข้อมูลแล้วควรเปิดหน้าจริงเพื่อตรวจผลเสมอ", STYLES["body"]))
    story.extend(img_flow("signin", "หน้าล็อกอิน Admin จากระบบจริงล่าสุด", max_h=2.55 * inch))
    story.extend(table_flow(MENU_ROWS, [1.25, 2.55, 2.65]))

    story.append(p("2. ความหมายของปุ่มสำคัญ", STYLES["h1"]))
    story.append(bullet_list([
        "Refresh: โหลดข้อมูลล่าสุดจาก Supabase อีกครั้ง ใช้หลังมีการแก้ไขข้อมูลหรือสงสัยว่าหน้าจอยังค้าง state เก่า",
        "Last synced: เวลาล่าสุดที่ระบบโหลดข้อมูลสำเร็จ ช่วยยืนยันว่าข้อมูลบนหน้าจอเป็นชุดล่าสุด",
        "Save draft: บันทึกเป็นฉบับร่างสำหรับ Products/Articles โดยยังไม่จำเป็นต้องเผยแพร่",
        "Publish: เผยแพร่ Products/Articles ให้พร้อมแสดงหน้าบ้าน",
        "Save changes: บันทึกเมนูที่ไม่มีระบบ draft/publish",
        "Discard changes: ยกเลิกการแก้ไขที่ยังไม่ได้บันทึก",
        "Delete: ลบรายการออกจากระบบ ควรใช้ด้วยความระมัดระวัง",
    ], STYLES["body"]))
    story.extend(img_flow("actions_product", "ชุดปุ่ม Save draft / Publish / Delete สำหรับสินค้าและบทความ", max_h=1.2 * inch))
    story.extend(img_flow("actions_settings", "ชุดปุ่ม Save changes / Discard changes / Delete สำหรับเมนูทั่วไป", max_h=1.65 * inch))

    story.append(p("3. คู่มือรายเมนู", STYLES["h1"]))
    for title, intro, bullets, note in SECTIONS:
        story.append(p(title, STYLES["h2"]))
        story.append(p(intro, STYLES["body"]))
        story.append(bullet_list(bullets, STYLES["body"]))
        story.append(p(f"<b>ข้อแนะนำ:</b> {note}", STYLES["note"]))
        if title == "All Products":
            story.extend(img_flow("product_editor", "ตัวอย่างช่องคำอธิบายสินค้าในระบบหลังบ้าน", max_h=1.25 * inch))
            story.extend(img_flow("product_public", "ตัวอย่างข้อมูลสินค้าที่แสดงในหน้าบ้านหลังอัปเดต", max_h=3.35 * inch))
        if title == "Articles":
            story.extend(img_flow("rich_text", "Rich Text editor มี scroll bar เพื่อรองรับเนื้อหายาว", max_h=2.85 * inch))
        if title == "Industries":
            story.extend(img_flow("industry_image_hidden", "ตัวอย่างส่วน Image URL เดิมที่ถูกซ่อนเพื่อไม่ให้ admin สับสน", max_h=1.9 * inch))

    story.append(PageBreak())
    story.append(p("4. Workflow แนะนำสำหรับลูกค้า", STYLES["h1"]))
    story.extend(table_flow([
        ["ขั้นตอน", "สิ่งที่ต้องทำ", "ผลลัพธ์ที่ควรตรวจ"],
        ["1", "เลือกเมนูและค้นหารายการ", "รายการที่เลือกตรงกับหน้าที่ต้องการแก้"],
        ["2", "แก้ข้อมูลในฟอร์ม", "Preview ด้านขวาอัปเดตตามข้อมูลใหม่"],
        ["3", "ตรวจ SEO/รูป/Slug", "ไม่มีข้อความซ้ำ รูปโหลดได้ URL ถูกต้อง"],
        ["4", "Save draft หรือ Save changes", "ไม่มี error และ Last synced อัปเดตหลัง Refresh"],
        ["5", "Publish ถ้าเป็น Products/Articles", "หน้าบ้านแสดงข้อมูลตามที่ตั้งใจ"],
        ["6", "เปิดหน้าจริงตรวจซ้ำ", "ไม่มี layout เพี้ยน รูปแตก หรือ link ผิด"],
    ], [0.65, 2.95, 2.85]))

    story.append(p("5. Checklist ก่อนเผยแพร่สินค้า/บทความ", STYLES["h1"]))
    story.append(bullet_list([
        "ชื่อเรื่องหรือชื่อสินค้าชัดเจนและสะกดถูก",
        "Slug เป็นภาษาอังกฤษ ตัวเล็ก ใช้ hyphen และไม่ซ้ำ",
        "รูปหลักคมชัด ไม่บีบผิดสัดส่วน และไม่มีข้อมูลลับ",
        "Description Text สรุปใจความ 1-3 ประโยค",
        "Rich Text มี heading, bullet, paragraph และ link ที่อ่านง่าย",
        "SEO Title และ Meta Description ไม่ยาวเกินไปและตรง keyword หลัก",
        "Preview / Google Search Preview ดูเรียบร้อย",
        "เปิด URL หน้าจริงหลังบันทึกเพื่อตรวจผล",
    ], STYLES["body"]))

    story.append(p("6. Troubleshooting เบื้องต้น", STYLES["h1"]))
    story.extend(table_flow([
        ["ปัญหา", "สาเหตุที่พบบ่อย", "วิธีแก้"],
        ["Save ไม่ผ่าน", "Session หมดอายุ, JSON ผิดรูปแบบ, field จำเป็นว่าง", "Refresh, login ใหม่, ตรวจ error และ JSON"],
        ["รูปไม่แสดง", "URL ผิด, ไฟล์ไม่ public, ขนาดใหญ่เกินไป", "Upload ใหม่หรือใช้ URL ที่ browser เปิดได้"],
        ["หน้าบ้านไม่อัปเดต", "ยังไม่ได้ Publish, cache หรือยังไม่ Refresh", "กด Publish/Save, Refresh admin และ reload หน้าบ้าน"],
        ["Slug ซ้ำ", "ใช้ URL ซ้ำกับรายการเดิม", "เปลี่ยน slug ให้เฉพาะเจาะจง"],
        ["JSON error", "ลืม comma/quote หรือโครงสร้างผิด", "ตรวจด้วย JSON validator ก่อนบันทึก"],
        ["LINE ไม่แจ้งเตือน", "Token ผิด, quota, endpoint ปิด", "ตรวจ LINE Settings และทดสอบอีกครั้ง"],
    ], [1.25, 2.4, 2.8]))

    story.append(
        KeepTogether([
            p("ข้อควรระวังด้านข้อมูลลับ", STYLES["h1"]),
            bullet_list([
        "ห้ามส่งต่อ LINE token, Supabase key, tracking ID แบบเต็ม หรือรหัสผ่าน admin ในเอกสารสาธารณะ",
        "ภาพหน้าจอสำหรับคู่มือควร crop เฉพาะ UI ที่ต้องอธิบาย",
        "เมนู LINE Settings และ Tracking Settings ควรจำกัดให้ผู้ดูแลระบบหลักหรือทีมเทคนิคแก้ไข",
        "ก่อนส่งคู่มือให้ลูกค้า ควรตรวจอีกครั้งว่าไม่มีค่า secret หรือข้อมูลส่วนตัวติดอยู่ในภาพ",
            ], STYLES["body"]),
        ])
    )

    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def make_styles():
    base = getSampleStyleSheet()
    return {
        "cover_title": ParagraphStyle("cover_title", parent=base["Title"], fontName="Tahoma-Bold", fontSize=25, leading=34, textColor=colors.HexColor("#001B44"), alignment=TA_CENTER, spaceAfter=4),
        "cover_subtitle": ParagraphStyle("cover_subtitle", parent=base["Title"], fontName="Tahoma-Bold", fontSize=15, leading=20, textColor=colors.HexColor("#008AF5"), alignment=TA_CENTER, spaceAfter=8),
        "kicker_center": ParagraphStyle("kicker_center", parent=base["Normal"], fontName="Tahoma-Bold", fontSize=9, leading=12, textColor=colors.HexColor("#008AF5"), alignment=TA_CENTER, spaceAfter=10),
        "lead_center": ParagraphStyle("lead_center", parent=base["Normal"], fontName="Tahoma", fontSize=10.2, leading=16, textColor=colors.HexColor("#5B677A"), alignment=TA_CENTER),
        "muted_center": ParagraphStyle("muted_center", parent=base["Normal"], fontName="Tahoma", fontSize=8.5, leading=12, textColor=colors.HexColor("#5B677A"), alignment=TA_CENTER),
        "h1": ParagraphStyle("h1", parent=base["Heading1"], fontName="Tahoma-Bold", fontSize=15.5, leading=22, textColor=colors.HexColor("#001B44"), spaceBefore=10, spaceAfter=8),
        "h2": ParagraphStyle("h2", parent=base["Heading2"], fontName="Tahoma-Bold", fontSize=12.6, leading=18, textColor=colors.HexColor("#008AF5"), spaceBefore=8, spaceAfter=5),
        "body": ParagraphStyle("body", parent=base["BodyText"], fontName="Tahoma", fontSize=9.45, leading=15, textColor=colors.HexColor("#001B44"), spaceAfter=6),
        "note": ParagraphStyle("note", parent=base["BodyText"], fontName="Tahoma", fontSize=8.7, leading=13, textColor=colors.HexColor("#5B677A"), backColor=colors.HexColor("#EEF7FF"), borderColor=colors.HexColor("#CDEBFF"), borderWidth=0.35, borderPadding=6, spaceAfter=8),
        "caption": ParagraphStyle("caption", parent=base["BodyText"], fontName="Tahoma", fontSize=8, leading=11, textColor=colors.HexColor("#5B677A"), alignment=TA_CENTER, spaceAfter=6),
        "table_header": ParagraphStyle("table_header", parent=base["BodyText"], fontName="Tahoma-Bold", fontSize=8.4, leading=12, textColor=colors.HexColor("#001B44"), alignment=TA_LEFT),
        "table_cell": ParagraphStyle("table_cell", parent=base["BodyText"], fontName="Tahoma", fontSize=8.15, leading=12, textColor=colors.HexColor("#001B44"), alignment=TA_LEFT),
    }


def prepare_images() -> None:
    source = TEMP / "codex-clipboard-cb661278-1acf-4479-9601-8e251ffa02d8.png"
    target = SCREENSHOTS["industry_image_hidden"]
    if source.exists():
        img = PILImage.open(source).convert("RGB")
        # Hide the long public storage URL so the manual focuses on the admin behavior.
        from PIL import ImageDraw, ImageFont

        draw = ImageDraw.Draw(img)
        draw.rounded_rectangle((18, 38, img.width - 18, 76), radius=14, fill=(255, 255, 255), outline=(218, 230, 242), width=2)
        font_path = r"C:\Windows\Fonts\tahoma.ttf"
        font = ImageFont.truetype(font_path, 18)
        draw.text((34, 47), "Image URL field is hidden in the updated admin UI", fill=(0, 27, 68), font=font)
        img.save(target)


if __name__ == "__main__":
    global STYLES
    register_fonts()
    STYLES = make_styles()
    prepare_images()
    build_docx()
    build_pdf()
    print(DOCX_PATH)
    print(PDF_PATH)
